import customtkinter as ctk
from tkinter import filedialog, messagebox
import openpyxl
import os
import sys
import threading

from autofix_core import (
    DEFAULT_TITLE,
    NEEDED,
    WIDTHS,
    ALIGNMENTS,
    ROWS_PER_GROUP,
    find_cols,
    save_as_excel,
)

ctk.set_appearance_mode("light")
ctk.set_default_color_theme("green")


class App(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("تعديل ملف الإكسيل الأسبوعي")
        self.geometry("850x700")
        self.minsize(700, 600)
        self.resizable(True, True)
        self.filepath = ""
        self.info = None
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(4, weight=1)

        # Header
        f = ctk.CTkFrame(self, corner_radius=0, fg_color=("gray90", "gray20"))
        f.grid(row=0, column=0, sticky="ew")
        f.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            f,
            text="📊 تعديل ملف الإكسيل الأسبوعي",
            font=ctk.CTkFont(size=24, weight="bold"),
            text_color=("green", "green"),
        ).grid(row=0, column=0, pady=(14, 2))
        ctk.CTkLabel(
            f,
            text="شيل الزايد، رتب حسب النوع، وحط صف عنوان",
            font=ctk.CTkFont(size=13),
            text_color=("gray30", "gray70"),
        ).grid(row=1, column=0, pady=(0, 10))

        # File
        f1 = ctk.CTkFrame(self)
        f1.grid(row=1, column=0, sticky="ew", padx=18, pady=(10, 4))
        f1.grid_columnconfigure(1, weight=1)
        ctk.CTkLabel(f1, text="الملف:", font=ctk.CTkFont(size=13, weight="bold")).grid(
            row=0, column=0, padx=(8, 2), pady=8
        )
        self.entry_file = ctk.CTkEntry(f1, placeholder_text="اضغط تصفح واختار الملف...")
        self.entry_file.grid(row=0, column=1, sticky="ew", padx=4, pady=8)
        ctk.CTkButton(
            f1, text="🔍 تصفح", width=90, command=self.browse, font=ctk.CTkFont(size=12)
        ).grid(row=0, column=2, padx=(4, 8), pady=8)

        # Title
        f2 = ctk.CTkFrame(self)
        f2.grid(row=2, column=0, sticky="ew", padx=18, pady=(4, 4))
        f2.grid_columnconfigure(1, weight=1)
        ctk.CTkLabel(
            f2, text="صف العنوان:", font=ctk.CTkFont(size=13, weight="bold")
        ).grid(row=0, column=0, padx=(8, 2), pady=8)
        self.entry_title = ctk.CTkEntry(f2, font=ctk.CTkFont(size=14))
        self.entry_title.insert(0, DEFAULT_TITLE)
        self.entry_title.grid(row=0, column=1, sticky="ew", padx=4, pady=8)

        # Status
        self.lbl_status = ctk.CTkLabel(
            self,
            text="🔘 في انتظار اختيار ملف...",
            font=ctk.CTkFont(size=12),
            anchor="w",
        )
        self.lbl_status.grid(row=3, column=0, sticky="ew", padx=22, pady=(4, 0))

        # Preview
        f3 = ctk.CTkFrame(self)
        f3.grid(row=4, column=0, sticky="nsew", padx=18, pady=(4, 8))
        f3.grid_rowconfigure(0, weight=1)
        f3.grid_columnconfigure(0, weight=1)
        tv = ctk.CTkTabview(f3)
        tv.grid(row=0, column=0, sticky="nsew")
        self.txt_b = ctk.CTkTextbox(
            tv.add("📂 قبل التعديل"),
            wrap="none",
            font=ctk.CTkFont(family="Consolas", size=11),
        )
        self.txt_b.pack(fill="both", expand=True, padx=4, pady=4)
        self.txt_a = ctk.CTkTextbox(
            tv.add("📋 بعد التعديل"),
            wrap="none",
            font=ctk.CTkFont(family="Consolas", size=11),
        )
        self.txt_a.pack(fill="both", expand=True, padx=4, pady=4)

        # Buttons
        f4 = ctk.CTkFrame(self, fg_color="transparent")
        f4.grid(row=5, column=0, sticky="ew", padx=18, pady=(0, 14))
        f4.grid_columnconfigure(0, weight=2)
        f4.grid_columnconfigure(1, weight=1)
        self.btn_go = ctk.CTkButton(
            f4,
            text="⚡ تنفيذ التعديل",
            font=ctk.CTkFont(size=15, weight="bold"),
            height=44,
            command=self.run,
            state="disabled",
            fg_color="#2e7d32",
            hover_color="#1b5e20",
        )
        self.btn_go.grid(row=0, column=0, padx=(0, 4), sticky="ew")
        ctk.CTkButton(
            f4,
            text="❌ خروج",
            font=ctk.CTkFont(size=14),
            height=44,
            command=self.quit,
            fg_color="#b71c1c",
            hover_color="#7f0000",
        ).grid(row=0, column=1, padx=(4, 0), sticky="ew")

    def browse(self):
        app_dir = os.path.dirname(os.path.abspath(sys.argv[0]))
        p = filedialog.askopenfilename(
            title="اختر ملف الإكسيل",
            initialdir=app_dir,
            filetypes=[("Excel files", "*.xlsx;*.xls"), ("All files", "*.*")],
        )
        if not p:
            return
        self.filepath = p
        self.entry_file.delete(0, "end")
        self.entry_file.insert(0, p)
        self.lbl_status.configure(text="🔍 جاري تحليل الملف...")
        self.btn_go.configure(state="disabled")
        self.update()
        threading.Thread(target=self.analyze, daemon=True).start()

    def analyze(self):
        try:
            wb = openpyxl.load_workbook(self.filepath, data_only=True)
            ws = wb.active
            hr = 1
            for r in range(1, min(ws.max_row + 1, 10)):
                for c in range(1, ws.max_column + 1):
                    v = ws.cell(r, c).value
                    if v and str(v).strip() == "CaseNumber":
                        hr = r
                        break
                else:
                    continue
                break

            hdrs = []
            for c in range(1, ws.max_column + 1):
                v = ws.cell(hr, c).value
                hdrs.append(str(v) if v else f"Col{c}")

            cols_map = {}
            for c in range(1, ws.max_column + 1):
                v = ws.cell(hr, c).value
                if v and str(v).strip() in NEEDED:
                    cols_map[str(v).strip()] = c

            found_cols = [cols_map.get(h) for h in NEEDED]

            sample = []
            for r in range(hr + 1, min(ws.max_row + 1, hr + 6)):
                row_v = []
                for c in range(1, ws.max_column + 1):
                    v = ws.cell(r, c).value
                    row_v.append(str(v)[:40] if v else "")
                sample.append(row_v)

            self.info = {
                "rows": ws.max_row,
                "cols": ws.max_column,
                "header_row": hr,
                "headers": hdrs,
                "found_cols": found_cols,
                "cols_map": cols_map,
                "sample": sample,
            }
            self.after(0, self.show_preview)
        except Exception as e:
            self.after(0, lambda: self.lbl_status.configure(text=f"❌ خطأ: {e}"))

    def show_preview(self):
        info = self.info
        title = self.entry_title.get().strip() or DEFAULT_TITLE
        cols = info["found_cols"]
        missing = [h for h, c in zip(NEEDED, cols) if not c]
        found = [h for h, c in zip(NEEDED, cols) if c]

        # Before
        self.txt_b.delete("0.0", "end")
        lines = [
            f"   📄 {os.path.basename(self.filepath)}",
            f"   📏 {info['rows']} صف x {info['cols']} عمود",
            f"   🏷️  صف الهيدر: {info['header_row']}",
        ]
        if missing:
            lines += ["", "   ❌ أعمدة مش موجودة:", f"      {', '.join(missing)}"]
        lines += ["", "   " + "─" * 70, "   HEADERS:"]
        for i, h in enumerate(info["headers"], 1):
            ok = h in NEEDED
            lines.append(f"   {'✅' if ok else '❌'} Col {i:2d}: {h}")
        lines += ["", "   " + "─" * 70, "   البيانات (أول 5):"]
        for ri, rv in enumerate(info["sample"]):
            parts = []
            for i, v in enumerate(rv):
                parts.append(f"{'✅' if info['headers'][i] in NEEDED else '❌'}{v}")
            lines.append(f"   صف {info['header_row'] + 1 + ri}: " + " | ".join(parts))
        self.txt_b.insert("0.0", "\n".join(lines))

        # After
        self.txt_a.delete("0.0", "end")
        kept = []
        for rv in info["sample"]:
            kept.append(
                [
                    rv[info["headers"].index(NEEDED[i])]
                    if NEEDED[i] in info["headers"]
                    else ""
                    for i in range(len(NEEDED))
                ]
            )
        if len(kept) > 0 and not missing:
            kept.sort(key=lambda x: str(x[4] or ""))

        lines2 = [
            f"   ✅ النتيجة: {len(info['sample'])} صف x {len(NEEDED)} أعمدة",
            f"   🔄 مرتب حسب: {NEEDED[4]}",
            f"   🏷️  صف العنوان: {title}",
        ]
        if missing:
            lines2 += ["", f"   ❌ مش هينفع! الأعمدة المفقودة: {', '.join(missing)}"]
        else:
            lines2 += [
                "",
                "   " + "─" * 70,
                f"   صف 1: {title}",
                f"   صف 2: {' | '.join(NEEDED)}",
                "   " + "─" * 70,
                "   البيانات (مرتبة):",
            ]
            for ri, rd in enumerate(kept):
                lines2.append(f"   صف {3 + ri}: {' | '.join(str(v)[:35] for v in rd)}")
            lines2 += ["", "   ✅ جاهز للتنفيذ!"]
        self.txt_a.insert("0.0", "\n".join(lines2))

        if missing:
            self.lbl_status.configure(text=f"❌ أعمدة مفقودة: {', '.join(missing)}")
            self.btn_go.configure(state="disabled")
        else:
            self.lbl_status.configure(text="✅ تم التحليل. اضغط 'تنفيذ التعديل'.")
            self.btn_go.configure(state="normal")

    def run(self):
        if not self.filepath or not os.path.exists(self.filepath):
            messagebox.showerror("خطأ", "الملف مش موجود!")
            return
        title = self.entry_title.get().strip() or DEFAULT_TITLE
        out_dir = os.path.dirname(self.filepath)
        self.btn_go.configure(state="disabled", text="⏳ جاري...")
        self.lbl_status.configure(text="⏳ جاري التعديل...")
        self.update()

        def go():
            try:
                path, cnt = save_as_excel(
                    self.filepath,
                    os.path.join(os.path.splitext(self.filepath)[0] + "_modified.xlsx"),
                    title,
                )
                self.after(0, lambda: self.done(path, cnt))
            except Exception as e:
                self.after(0, lambda: self.err(str(e)))

        threading.Thread(target=go, daemon=True).start()

    def done(self, path, cnt):
        self.btn_go.configure(state="normal", text="⚡ تنفيذ التعديل")
        self.last_excel_path = path
        msg = f"تم تعديل {cnt} صف.\n📁 {os.path.basename(path)}"
        self.lbl_status.configure(text=f"✅ تم! الملف: {path}")
        if messagebox.askyesno("✅ تم", msg + "\n\nفتح المجلد؟"):
            os.startfile(os.path.dirname(path))

    def err(self, msg):
        self.btn_go.configure(state="normal", text="⚡ تنفيذ التعديل")
        self.lbl_status.configure(text=f"❌ خطأ: {msg}")
        messagebox.showerror("خطأ", msg)

    def quit(self):
        if messagebox.askokcancel("خروج", "خروج من البرنامج؟"):
            self.destroy()


if __name__ == "__main__":
    App().mainloop()
